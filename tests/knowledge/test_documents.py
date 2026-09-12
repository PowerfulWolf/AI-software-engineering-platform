"""Company document imports are immutable, bounded and directly usable as context."""

from __future__ import annotations

import hashlib
import json
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from ai_software_engineer.company_workspace import (
    CompanyWorkspace,
    discover_company_workspaces,
)
from ai_software_engineer.knowledge_documents import (
    CompanyKnowledgeDocumentStore,
    KnowledgeDocumentError,
)


def _company(tmp_path: Path) -> CompanyWorkspace:
    return CompanyWorkspace.initialize(
        tmp_path / "platform", company_id="company_test", name="Test company"
    )


def _docx(text: str) -> bytes:
    document = Document()
    document.add_heading("Engineering", level=1)
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        (
            f"<< /Length {len(escaped) + 33} >>\nstream\n"
            f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET\nendstream"
        ).encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(value)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def test_company_catalog_and_markdown_import_are_stable_and_context_ready(
    tmp_path: Path,
) -> None:
    company = _company(tmp_path)
    store = CompanyKnowledgeDocumentStore(company)

    first = store.import_document(filename="team-guide.md", content=b"# Team\n\nReview first.\n")
    replay = store.import_document(filename="renamed.md", content=b"# Team\n\nReview first.\n")

    assert replay == first
    assert store.list() == (first,)
    assert discover_company_workspaces(tmp_path / "platform") == (company,)
    sources = company.knowledge_sources((first.normalized_relative_path,))
    assert sources[0].content == "# Team\n\nReview first.\n"
    assert first.source_sha256 != "0" * 64
    assert (
        (store.root / first.document_id / first.source_relative_path)
        .read_bytes()
        .startswith(b"# Team")
    )


@pytest.mark.parametrize(
    ("filename", "content", "expected"),
    [
        ("guide.txt", b"Release checklist", "Release checklist"),
        ("guide.docx", _docx("No self approval."), "No self approval."),
        ("guide.pdf", _pdf("Evidence required"), "Evidence required"),
    ],
)
def test_supported_documents_are_normalized(
    tmp_path: Path, filename: str, content: bytes, expected: str
) -> None:
    store = CompanyKnowledgeDocumentStore(_company(tmp_path))

    manifest = store.import_document(filename=filename, content=content)

    normalized = (store.root / manifest.document_id / "content.md").read_text()
    assert normalized.startswith("# guide")
    assert expected in normalized


@pytest.mark.parametrize("filename", ["../guide.md", ".env", "guide.exe", "secrets.txt"])
def test_unsafe_document_names_are_rejected(tmp_path: Path, filename: str) -> None:
    with pytest.raises(KnowledgeDocumentError, match="unsafe"):
        CompanyKnowledgeDocumentStore(_company(tmp_path)).import_document(
            filename=filename, content=b"content"
        )


def test_empty_invalid_and_tampered_documents_fail_closed(tmp_path: Path) -> None:
    store = CompanyKnowledgeDocumentStore(_company(tmp_path))
    with pytest.raises(KnowledgeDocumentError, match="empty"):
        store.import_document(filename="empty.txt", content=b"")
    with pytest.raises(KnowledgeDocumentError, match="decoded"):
        store.import_document(filename="broken.docx", content=b"not a document")
    manifest = store.import_document(filename="guide.txt", content=b"Trusted text")
    (store.root / manifest.document_id / "content.md").write_text("changed")
    with pytest.raises(KnowledgeDocumentError, match="integrity"):
        store.list()


def test_source_upload_limit_is_enforced_before_extraction(tmp_path: Path) -> None:
    store = CompanyKnowledgeDocumentStore(_company(tmp_path))

    with pytest.raises(KnowledgeDocumentError, match="upload limit"):
        store.import_document(filename="large.txt", content=b"x" * 10_000_001)


def test_oversized_normalized_content_and_resealed_path_escape_are_rejected(
    tmp_path: Path,
) -> None:
    store = CompanyKnowledgeDocumentStore(_company(tmp_path))
    with pytest.raises(KnowledgeDocumentError, match="context limit"):
        store.import_document(filename="large.txt", content=b"x" * 256_001)
    manifest = store.import_document(filename="guide.md", content=b"# Guide\n")
    manifest_path = store.root / manifest.document_id / "manifest.json"
    payload = json.loads(manifest_path.read_text())
    payload["source_relative_path"] = "../outside.md"
    canonical = json.dumps(
        {key: value for key, value in payload.items() if key != "manifest_sha256"},
        sort_keys=True,
        separators=(",", ":"),
    ).encode()
    payload["manifest_sha256"] = hashlib.sha256(canonical).hexdigest()
    manifest_path.write_text(json.dumps(payload))

    with pytest.raises(KnowledgeDocumentError, match="manifest is invalid"):
        store.list()


def test_resealed_source_name_suffix_drift_is_rejected(tmp_path: Path) -> None:
    store = CompanyKnowledgeDocumentStore(_company(tmp_path))
    manifest = store.import_document(filename="guide.md", content=b"# Guide\n")
    manifest_path = store.root / manifest.document_id / "manifest.json"
    payload = json.loads(manifest_path.read_text())
    payload["source_name"] = "guide.pdf"
    canonical = json.dumps(
        {key: value for key, value in payload.items() if key != "manifest_sha256"},
        sort_keys=True,
        separators=(",", ":"),
    ).encode()
    payload["manifest_sha256"] = hashlib.sha256(canonical).hexdigest()
    manifest_path.write_text(json.dumps(payload))

    with pytest.raises(KnowledgeDocumentError, match="manifest is invalid"):
        store.list()
