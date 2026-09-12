"""Immutable, content-addressed Company knowledge documents."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import stat
import tempfile
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path, PurePath
from typing import Annotated, Literal, Self
from zipfile import BadZipFile, ZipFile

from docx import Document
from pydantic import AwareDatetime, Field, StringConstraints, model_validator
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from ai_software_engineer.company_workspace import (
    MAX_COMPANY_KNOWLEDGE_DOCUMENT_BYTES,
    MAX_COMPANY_KNOWLEDGE_SOURCE_BYTES,
    CompanyId,
    CompanyWorkspace,
    Digest,
    validate_knowledge_path,
)
from ai_software_engineer.domain.model import DomainModel

KnowledgeDocumentId = Annotated[
    str, StringConstraints(pattern=r"^knowledge_document_[a-f0-9]{32}$")
]
SourceName = Annotated[str, StringConstraints(min_length=1, max_length=200)]
KnowledgeMediaType = Literal["text/markdown", "text/plain", "application/pdf", "application/docx"]
_ALLOWED_SUFFIXES = {".md", ".txt", ".pdf", ".docx"}


class KnowledgeDocumentError(RuntimeError):
    """Raised when an uploaded document cannot be safely imported."""


class KnowledgeDocumentManifest(DomainModel):
    schema_version: Literal["v0.1"] = "v0.1"
    document_id: KnowledgeDocumentId
    company_id: CompanyId
    source_name: SourceName
    media_type: KnowledgeMediaType
    source_relative_path: SourceName
    normalized_relative_path: Annotated[str, StringConstraints(min_length=1, max_length=300)]
    source_bytes: Annotated[int, Field(ge=1, le=MAX_COMPANY_KNOWLEDGE_SOURCE_BYTES)]
    normalized_bytes: Annotated[int, Field(ge=1, le=MAX_COMPANY_KNOWLEDGE_DOCUMENT_BYTES)]
    source_sha256: Digest
    normalized_sha256: Digest
    imported_at: AwareDatetime
    manifest_sha256: Digest

    @model_validator(mode="after")
    def validate_paths_and_identity(self) -> Self:
        suffix = Path(self.source_relative_path).suffix.lower()
        try:
            _validate_source_name(self.source_name)
        except KnowledgeDocumentError as error:
            raise ValueError("knowledge document source name is invalid") from error
        if (
            self.source_relative_path != f"source{suffix}"
            or suffix not in _ALLOWED_SUFFIXES
            or Path(self.source_name).suffix.lower() != suffix
            or self.media_type != _media_type(suffix)
            or self.normalized_relative_path != f"documents/{self.document_id}/content.md"
            or self.document_id != "knowledge_document_" + self.source_sha256[:32]
        ):
            raise ValueError("knowledge document manifest paths or identity are invalid")
        validate_knowledge_path(self.normalized_relative_path)
        return self

    def recompute_digest(self) -> str:
        payload = self.model_dump(mode="json", exclude={"manifest_sha256"})
        return hashlib.sha256(
            json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
        ).hexdigest()

    def validate_integrity(self) -> None:
        if self.manifest_sha256 != self.recompute_digest():
            raise KnowledgeDocumentError("knowledge document manifest digest mismatch")


@dataclass(frozen=True)
class CompanyKnowledgeDocumentStore:
    company: CompanyWorkspace

    @property
    def root(self) -> Path:
        return self.company.root / "knowledge" / "documents"

    def list(self) -> tuple[KnowledgeDocumentManifest, ...]:
        self.company.validate_current()
        if not self.root.exists():
            return ()
        _reject_symlink(self.root)
        manifests: list[KnowledgeDocumentManifest] = []
        for directory in self.root.iterdir():
            if directory.is_symlink():
                raise KnowledgeDocumentError("knowledge document storage cannot traverse a symlink")
            if directory.is_dir() and not directory.name.startswith(".knowledge-"):
                manifests.append(self._read(directory))
        return tuple(sorted(manifests, key=lambda item: (item.imported_at, item.document_id)))

    def import_document(
        self,
        *,
        filename: str,
        content: bytes,
        imported_at: datetime | None = None,
    ) -> KnowledgeDocumentManifest:
        self.company.validate_current()
        source_name = _validate_source_name(filename)
        if not content:
            raise KnowledgeDocumentError("knowledge document is empty")
        if len(content) > MAX_COMPANY_KNOWLEDGE_SOURCE_BYTES:
            raise KnowledgeDocumentError("knowledge document exceeds upload limit")
        suffix = Path(source_name).suffix.lower()
        source_sha = hashlib.sha256(content).hexdigest()
        document_id = "knowledge_document_" + source_sha[:32]
        normalized = _normalize(source_name, content)
        normalized_bytes = normalized.encode("utf-8")
        if not normalized.strip():
            raise KnowledgeDocumentError("knowledge document contains no extractable text")
        if len(normalized_bytes) > MAX_COMPANY_KNOWLEDGE_DOCUMENT_BYTES:
            raise KnowledgeDocumentError("normalized knowledge document exceeds context limit")
        source_relative = f"source{suffix}"
        normalized_relative = f"documents/{document_id}/content.md"
        validate_knowledge_path(normalized_relative)
        media_type = _media_type(suffix)
        timestamp = imported_at or datetime.now(UTC)
        if timestamp.tzinfo is None:
            raise KnowledgeDocumentError("knowledge import timestamp must include a timezone")
        provisional = KnowledgeDocumentManifest(
            document_id=document_id,
            company_id=self.company.manifest.company_id,
            source_name=source_name,
            media_type=media_type,
            source_relative_path=source_relative,
            normalized_relative_path=normalized_relative,
            source_bytes=len(content),
            normalized_bytes=len(normalized_bytes),
            source_sha256=source_sha,
            normalized_sha256=hashlib.sha256(normalized_bytes).hexdigest(),
            imported_at=timestamp,
            manifest_sha256="0" * 64,
        )
        manifest = provisional.model_copy(
            update={"manifest_sha256": provisional.recompute_digest()}
        )
        target = self.root / document_id
        if target.exists():
            existing = self._read(target)
            if existing.source_sha256 != source_sha:
                raise KnowledgeDocumentError("knowledge document identity collision")
            return existing
        self.root.mkdir(parents=True, exist_ok=True)
        _reject_symlink(self.root)
        staging = Path(tempfile.mkdtemp(prefix=".knowledge-", dir=self.root))
        try:
            _write_new(staging / source_relative, content)
            _write_new(staging / "content.md", normalized_bytes)
            _write_new(staging / "manifest.json", manifest.model_dump_json(indent=2).encode())
            _sync_directory(staging)
            try:
                staging.rename(target)
            except OSError:
                if not target.exists():
                    raise
            _sync_directory(self.root)
        finally:
            if staging.exists():
                shutil.rmtree(staging)
        return self._read(target)

    def _read(self, directory: Path) -> KnowledgeDocumentManifest:
        _reject_symlink(directory)
        manifest_path = directory / "manifest.json"
        try:
            payload = _read_bounded(manifest_path, 32_000)
        except OSError as error:
            raise KnowledgeDocumentError("knowledge document record is incomplete") from error
        try:
            manifest = KnowledgeDocumentManifest.model_validate_json(payload)
        except ValueError as error:
            raise KnowledgeDocumentError("knowledge document manifest is invalid") from error
        manifest.validate_integrity()
        if (
            manifest.company_id != self.company.manifest.company_id
            or directory.name != manifest.document_id
        ):
            raise KnowledgeDocumentError("knowledge document identity mismatch")
        source = directory / manifest.source_relative_path
        normalized = directory / "content.md"
        source_bytes = _read_bounded(source, MAX_COMPANY_KNOWLEDGE_SOURCE_BYTES)
        normalized_bytes = _read_bounded(normalized, MAX_COMPANY_KNOWLEDGE_DOCUMENT_BYTES)
        if (
            len(source_bytes) != manifest.source_bytes
            or len(normalized_bytes) != manifest.normalized_bytes
            or hashlib.sha256(source_bytes).hexdigest() != manifest.source_sha256
            or hashlib.sha256(normalized_bytes).hexdigest() != manifest.normalized_sha256
        ):
            raise KnowledgeDocumentError("knowledge document content integrity mismatch")
        return manifest


def _validate_source_name(value: str) -> str:
    if (
        not value
        or len(value) > 200
        or PurePath(value).name != value
        or "/" in value
        or "\\" in value
        or value.startswith(".")
        or any(ord(character) < 32 for character in value)
        or Path(value).suffix.lower() not in _ALLOWED_SUFFIXES
        or Path(value).stem.lower() in {"credentials", "secrets", "id_rsa", "id_ed25519"}
    ):
        raise KnowledgeDocumentError("unsupported or unsafe knowledge document name")
    return value


def _media_type(suffix: str) -> KnowledgeMediaType:
    values: dict[str, KnowledgeMediaType] = {
        ".md": "text/markdown",
        ".txt": "text/plain",
        ".pdf": "application/pdf",
        ".docx": "application/docx",
    }
    return values[suffix]


def _normalize(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in {".md", ".txt"}:
            text = content.decode("utf-8")
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted and not reader.decrypt(""):
                raise KnowledgeDocumentError("encrypted PDF documents are not supported")
            if len(reader.pages) > 500:
                raise KnowledgeDocumentError("PDF document exceeds the page limit")
            pages = []
            for index, page in enumerate(reader.pages, start=1):
                extracted = page.extract_text() or ""
                if extracted.strip():
                    pages.append(f"## Page {index}\n\n{extracted.strip()}")
            text = "\n\n".join(pages)
        else:
            with ZipFile(BytesIO(content)) as archive:
                if (
                    len(archive.infolist()) > 2_000
                    or sum(item.file_size for item in archive.infolist()) > 25_000_000
                ):
                    raise KnowledgeDocumentError("DOCX document exceeds the expanded size limit")
            document = Document(BytesIO(content))
            parts = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            text = "\n\n".join(parts)
    except KnowledgeDocumentError:
        raise
    except (BadZipFile, OSError, PdfReadError, ValueError, UnicodeError) as error:
        raise KnowledgeDocumentError("knowledge document cannot be decoded") from error
    if suffix == ".md":
        return text.strip() + "\n"
    title = Path(filename).stem.replace("#", "").strip() or "Imported document"
    return f"# {title}\n\n{text.strip()}\n"


def _reject_symlink(path: Path) -> None:
    if any(item.is_symlink() for item in (path, *path.parents)):
        raise KnowledgeDocumentError("knowledge document storage cannot traverse a symlink")


def _write_new(path: Path, content: bytes) -> None:
    with path.open("xb") as stream:
        stream.write(content)
        stream.flush()
        os.fsync(stream.fileno())


def _read_bounded(path: Path, limit: int) -> bytes:
    _reject_symlink(path)
    try:
        descriptor = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
        with os.fdopen(descriptor, "rb") as stream:
            if not stat.S_ISREG(os.fstat(stream.fileno()).st_mode):
                raise KnowledgeDocumentError("knowledge document record is not a regular file")
            data = stream.read(limit + 1)
    except OSError as error:
        raise KnowledgeDocumentError("knowledge document record is incomplete") from error
    if len(data) > limit:
        raise KnowledgeDocumentError("knowledge document record exceeds its limit")
    return data


def _sync_directory(path: Path) -> None:
    descriptor = os.open(path, os.O_RDONLY)
    try:
        os.fsync(descriptor)
    finally:
        os.close(descriptor)


__all__ = [
    "CompanyKnowledgeDocumentStore",
    "KnowledgeDocumentError",
    "KnowledgeDocumentManifest",
]
